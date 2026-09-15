from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from model_workflow_common import load_config, validate_public_config


DOMAIN_ORDER = ["clean_common_domain", "canonical_full_domain"]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the model-comparison workbook from computed workflow results.")
    parser.add_argument("--config", default="config.yaml", help="Path to the YAML configuration file.")
    return parser.parse_args()


def metric_text(row: pd.Series, metric: str, digits: int = 3) -> str:
    mean = row.get(f"{metric}_mean")
    sd = row.get(f"{metric}_seed_SD")
    if pd.isna(mean):
        return ""
    if pd.isna(sd):
        return f"{float(mean):.{digits}f}"
    return f"{float(mean):.{digits}f} ± {float(sd):.{digits}f}"


def bootstrap_text(bootstrap: pd.DataFrame, domain: str, model: str, metric: str) -> str:
    rows = bootstrap.loc[
        (bootstrap["domain"] == domain)
        & (bootstrap["candidate_model"] == model)
        & (bootstrap["metric"] == metric)
    ]
    if rows.empty:
        return ""
    row = rows.iloc[0]
    return (
        f"{float(row['delta_point_estimate']):.3f} "
        f"[{float(row['ci_low_2.5']):.3f}, {float(row['ci_high_97.5']):.3f}]"
    )


def style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_cells in sheet.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(length + 2, 10), 42)


def main() -> None:
    args = parse_args()
    config = load_config(args.config)
    validate_public_config(config)
    results_dir = config.path("results_dir")
    summary = pd.read_csv(results_dir / "workflow_13_model_dual_domain_metrics.csv")
    bootstrap = pd.read_csv(results_dir / "workflow_ablation_bootstrap_vs_full.csv")
    model_order = [model["name"] for model in config.models]
    model_meta = {model["name"]: model for model in config.models}

    indexed = summary.set_index(["domain", "model"])
    workbook = Workbook()
    overview = workbook.active
    overview.title = "13-model overview"
    detail = workbook.create_sheet("Numeric detail")
    definitions = workbook.create_sheet("Definitions & sources")

    overview.append(
        [
            "Model",
            "Role",
            "Architecture / definition",
            "Tuning / fairness",
            "Clean RMSE",
            "Clean MAE",
            "Clean R²",
            "Clean Bias",
            "Clean r",
            "ΔRMSE vs Full [95% CI]",
            "Canonical RMSE",
            "Canonical MAE",
            "Canonical R²",
            "Canonical Bias",
            "Canonical r",
            "Runs",
        ]
    )

    for model in model_order:
        clean = indexed.loc[("clean_common_domain", model)]
        canonical = indexed.loc[("canonical_full_domain", model)]
        meta = model_meta[model]
        overview.append(
            [
                model,
                meta["role"],
                meta["architecture"],
                meta["tuning"],
                metric_text(clean, "RMSE"),
                metric_text(clean, "MAE"),
                metric_text(clean, "R2", 4),
                metric_text(clean, "Bias"),
                metric_text(clean, "Pearson_r", 4),
                "" if model == "Full canonical" else bootstrap_text(bootstrap, "clean_common_domain", model, "RMSE"),
                metric_text(canonical, "RMSE"),
                metric_text(canonical, "MAE"),
                metric_text(canonical, "R2", 4),
                metric_text(canonical, "Bias"),
                metric_text(canonical, "Pearson_r", 4),
                int(clean["runs"]),
            ]
        )

    detail.append(
        [
            "Domain",
            "Model",
            "Runs",
            "RMSE mean",
            "RMSE SD",
            "MAE mean",
            "MAE SD",
            "R² mean",
            "R² SD",
            "Bias mean",
            "Bias SD",
            "Pearson r mean",
            "Pearson r SD",
        ]
    )
    for domain in DOMAIN_ORDER:
        for model in model_order:
            row = indexed.loc[(domain, model)]
            detail.append(
                [
                    domain,
                    model,
                    int(row["runs"]),
                    row["RMSE_mean"],
                    row["RMSE_seed_SD"],
                    row["MAE_mean"],
                    row["MAE_seed_SD"],
                    row["R2_mean"],
                    row["R2_seed_SD"],
                    row["Bias_mean"],
                    row["Bias_seed_SD"],
                    row["Pearson_r_mean"],
                    row["Pearson_r_seed_SD"],
                ]
            )

    definitions.append(["Item", "Definition"])
    definitions.append(["Historical window", f"W{int(config.raw['experiment']['window_length'])}"])
    definitions.append(["Training years", "-".join(map(str, config.raw["experiment"]["train_years"]))])
    definitions.append(["Validation years", "-".join(map(str, config.raw["experiment"]["validation_years"]))])
    definitions.append(["Evaluation years", "-".join(map(str, config.raw["experiment"]["evaluation_years"]))])
    definitions.append(["Training seeds", ", ".join(map(str, config.seeds))])
    definitions.append(["Bias", "Mean(prediction - reference)"])
    definitions.append(["Seed variability", "Sample standard deviation across configured training seeds (ddof=1)"])
    definitions.append(
        [
            "Bootstrap",
            (
                f"Paired {config.raw['bootstrap']['cluster_variable']}-cluster bootstrap; "
                f"B={config.raw['bootstrap']['n_resamples']}; "
                f"seed={config.raw['bootstrap']['random_seed']}; candidate - Full canonical"
            ),
        ]
    )
    definitions.append(["No BiLSTM", "CNN -> Attention -> Regressor; recurrent/BiLSTM block removed"])

    for sheet in workbook.worksheets:
        style_sheet(sheet)

    output_path = config.path("workbook")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(f"Wrote {output_path}")


if __name__ == "__main__":
    main()


